import streamlit as st
import pandas as pd
from io import BytesIO
import openpyxl
import os

st.set_page_config(
    page_title="Feuille de Garde CIS CARSALADE", 
    page_icon="🚒", 
    layout="wide"
)

st.markdown("""
    <style>
    .main { background-color: #0e1117; }
    .header-box {
        background: linear-gradient(90deg, #b71c1c 0%, #d32f2f 100%);
        padding: 15px; border-radius: 8px; color: white;
        text-align: center; margin-bottom: 20px;
    }
    .header-title { font-size: 1.8rem; font-weight: 800; margin: 0; }
    .header-subtitle { font-size: 1rem; font-style: italic; margin: 5px 0 0 0; opacity: 0.9; }
    .section-title {
        color: #ff5252; font-size: 1.3rem; font-weight: 700;
        margin-top: 25px; margin-bottom: 10px;
        border-bottom: 2px solid #333; padding-bottom: 5px;
    }
    .vsav-title {
        color: #4fc3f7;
        font-size: 1.5rem; font-weight: 800;
        margin-top: 10px; margin-bottom: 10px;
        border-bottom: 3px solid #0288d1; padding-bottom: 5px;
    }
    div[data-baseweb="input"] > input { font-weight: bold; color: #ff5252; }
    .role-text { font-size: 1.2rem; font-weight: bold; text-align: right; padding-top: 35px; }
    </style>
""", unsafe_allow_html=True)

GOOGLE_SHEET_ID = "1WbCH8Q4r2rM2WL1f8KP2o7XaADi-1vjC"
KNOWN_BASES = ["VSAV", "VSRM", "VBAL", "VID", "VSMPM", "FPT", "EPC", "CCFM", "VFCDG"]

@st.cache_data(ttl=60)
def charger_donnees_depuis_gsheets(sheet_id):
    try:
        url_billet = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet=BILLET"
        df_brut = pd.read_csv(url_billet, header=None)
    except Exception:
        df_brut = pd.DataFrame()

    donnees = []
    fonctions_valides = ["CA", "COND", "EQ", "CE B1", "EQ B1", "CE B2", "EQ B2", "OBS", "COND/EQ"]
    mots_ignores = ["BILLET", "DE", "GARDE", "EQUIPE", "LUNDI", "MARDI", "MERCREDI", 
                    "JEUDI", "VENDREDI", "SAMEDI", "DIMANCHE", "JANVIER", "FEVRIER", 
                    "MARS", "AVRIL", "MAI", "JUIN", "JUILLET", "AOUT", "SEPTEMBRE", 
                    "OCTOBRE", "NOVEMBRE", "DECEMBRE", "CONSIGNES", "SPORT", "FMA",
                    "SPECIALITE", "SPECIALITÉ", "SPÉCIALITÉ", "SPÉCIALITE", "JOUR", "NUIT", "COMPETENCE"]
    
    INDICATIFS_AUTO = {
        "VSRM": "VSRM 10", "VFCDG": "VFCDG 88", "FPT": "FPT 11",
        "EPC": "EPC 13", "VID": "VID 34", "CCFM": "CCFM 29", "VSMPM": "VSMPM 02"
    }

    nom_engin_actuel = "GENERAL"
    vehicules_vus = {}

    if not df_brut.empty:
        for col_idx in range(df_brut.shape[1]):
            for row_idx in range(df_brut.shape[0]):
                valeur = str(df_brut.iloc[row_idx, col_idx]).strip()
                if pd.isna(df_brut.iloc[row_idx, col_idx]) or valeur.lower() == "nan" or valeur == "":
                    continue
                valeur_maj = valeur.upper()
                
                is_fonction = valeur_maj in fonctions_valides
                is_ignore = any(ignore in valeur_maj for ignore in mots_ignores)
                
                if is_fonction:
                    personnel = ""
                    if col_idx + 1 < df_brut.shape[1]:
                        p = str(df_brut.iloc[row_idx, col_idx + 1]).strip()
                        if p.lower() != "nan" and p != "" and p.upper() not in fonctions_valides:
                            personnel = p
                    donnees.append({"Agrès": nom_engin_actuel, "Fonction": valeur_maj, "Personnel": personnel})
                elif not is_ignore and len(valeur_maj) >= 2 and not is_fonction:
                    num_indicatif = ""
                    if col_idx + 1 < df_brut.shape[1]:
                        val_suiv = str(df_brut.iloc[row_idx, col_idx + 1]).strip()
                        if val_suiv.replace('.', '', 1).isdigit():
                            num_indicatif = str(int(float(val_suiv)))
                    elif row_idx + 1 < df_brut.shape[0]:
                        val_dessous = str(df_brut.iloc[row_idx + 1, col_idx]).strip()
                        if val_dessous.replace('.', '', 1).isdigit():
                            num_indicatif = str(int(float(val_dessous)))
                            
                    base_name = f"{valeur_maj} {num_indicatif}" if num_indicatif else valeur_maj
                    
                    if base_name == "VSAV":
                        if base_name in vehicules_vus:
                            vehicules_vus[base_name] += 1
                            nom_engin_actuel = "VSAV 17" if vehicules_vus[base_name] == 2 else f"VSAV {vehicules_vus[base_name]}"
                        else:
                            vehicules_vus[base_name] = 1
                            nom_engin_actuel = "VSAV 98"
                    elif base_name in INDICATIFS_AUTO:
                        nom_engin_actuel = INDICATIFS_AUTO[base_name]
                    else:
                        if base_name in vehicules_vus:
                            vehicules_vus[base_name] += 1
                            nom_engin_actuel = f"{base_name} {vehicules_vus[base_name]}"
                            if vehicules_vus[base_name] == 2:
                                for d in donnees:
                                    if d["Agrès"] == base_name: d["Agrès"] = f"{base_name} 1"
                        else:
                            vehicules_vus[base_name] = 1
                            nom_engin_actuel = base_name
                    
    df_garde = pd.DataFrame(donnees)

    liste_agents = []
    dict_agents = {}
    dict_specs = {}
    dict_comp_string = {} 

    try:
        url_effectif = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet=EFFECTIF"
        df_eff_brut = pd.read_csv(url_effectif, header=None)
        
        for r in range(1, len(df_eff_brut)):
            row = df_eff_brut.iloc[r]
            nom = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ""
            prenom = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
            grade = str(row.iloc[2]).strip() if pd.notna(row.iloc[2]) else ""
            
            if not nom or nom.lower() == "nan": continue
            nom_simple = f"{nom} {prenom}"
            
            specs_filtre = []
            for c in range(2, len(row)): 
                val = row.iloc[c]
                if pd.notna(val) and str(val).strip() != "" and str(val).lower() != "nan":
                    specs_filtre.append(str(val).strip().upper())
            
            specs_affichage = []
            for c in range(5, len(row)): 
                val = row.iloc[c]
                if pd.notna(val) and str(val).strip() != "" and str(val).lower() != "nan":
                    specs_affichage.append(str(val).strip().upper())
            
            details = []
            if grade and grade.lower() != 'nan':
                details.append(grade)
            if specs_affichage:
                details.extend(specs_affichage)
                
            label = f"{nom_simple} ({' - '.join(details)})" if details else nom_simple
                
            liste_agents.append(label)
            dict_agents[label] = nom_simple
            dict_specs[label] = specs_filtre 
            dict_comp_string[label] = " - ".join(details) 
            
        liste_agents = sorted(list(set(liste_agents)))
    except Exception:
        pass

    return df_garde, liste_agents, dict_agents, dict_specs, dict_comp_string

st.markdown("""
    <div class="header-box">
        <p class="header-title">🚒 CIS CARSALADE</p>
        <p class="header-subtitle">Feuille de Garde - Secours à Personne prioritaires</p>
    </div>
""", unsafe_allow_html=True)

try:
    df_garde, liste_agents, dict_agents, dict_specs, dict_comp_string = charger_donnees_depuis_gsheets(GOOGLE_SHEET_ID)
    
    if df_garde.empty:
        st.warning("⚠️ Impossible de lire l'onglet 'BILLET' de votre Google Sheets.")
    else:
        def get_options_filtrees(agres_nom, fonction, agent_actuel_str):
            options_valides = [""]
            for opt in liste_agents:
                specs_agent = dict_specs.get(opt, [])
                est_autorise = True
                specs_clean = [str(s).replace(" ", "").replace("-", "").upper() for s in specs_agent]
                
                if "VSAV" in agres_nom.upper() and fonction.upper() == "CA":
                    if not any(kw in specs_clean for kw in ["CA", "CA1E", "CATE", "CAVSAV"]):
                        est_autorise = False
                elif "FPT" in agres_nom.upper() and fonction.upper() == "CA":
                    if not any(kw in specs_clean for kw in ["CATE", "CAFPT"]):
                        est_autorise = False
                        
                if est_autorise or (agent_actuel_str.strip() != "" and agent_actuel_str.strip().lower() in opt.lower()):
                    options_valides.append(opt)
            return options_valides

        agres_uniques = df_garde['Agrès'].unique()
        export_data = [] 
        
        vsav_uniques = [a for a in agres_uniques if "VSAV" in a.upper()]
        autres_uniques = [a for a in agres_uniques if "VSAV" not in a.upper()]

        def afficher_bloc_engin(liste_agres, title_html=None):
            if not liste_agres: return
            if title_html: st.markdown(title_html, unsafe_allow_html=True)
            
            for i_veh in range(0, len(liste_agres), 3):
                cols_ligne = st.columns(3)
                batch = liste_agres[i_veh:i_veh+3]
                
                for idx_col, agres in enumerate(batch):
                    df_agres = df_garde[df_garde['Agrès'] == agres]
                    base_name_pure = next((b for b in KNOWN_BASES if b in agres), agres.split(" ")[0])
                    
                    with cols_ligne[idx_col]:
                        nouveau_nom_agres = st.text_input("🚚 Engin :", value=agres, key=f"edit_engin_{agres}")
                        veh_data = {'base_name': base_name_pure, 'custom_name': nouveau_nom_agres, 'roles': []}
                        
                        for i, row in df_agres.iterrows():
                            fonction = row['Fonction']
                            agent_actuel = row['Personnel'] 
                            
                            c_role, c_jour, c_nuit = st.columns([0.8, 2, 2])
                            with c_role:
                                st.markdown(f'<div class="role-text">{fonction}</div>', unsafe_allow_html=True)
                            with c_jour:
                                options_j = get_options_filtrees(nouveau_nom_agres, fonction, agent_actuel)
                                def_idx = next((idx for idx, opt in enumerate(options_j) if agent_actuel.strip().lower() in opt.lower() and agent_actuel.strip()!=""), 0)
                                jour_lbl = st.selectbox("☀️ Jour", options=options_j, index=def_idx, key=f"j_{agres}_{fonction}_{i}")
                            with c_nuit:
                                options_n = get_options_filtrees(nouveau_nom_agres, fonction, "")
                                nuit_lbl = st.selectbox("🌙 Nuit", options=options_n, index=0, key=f"n_{agres}_{fonction}_{i}")
                            
                            veh_data['roles'].append({
                                'fonction': fonction,
                                'jour_name': dict_agents.get(jour_lbl, ""),
                                'jour_comp': dict_comp_string.get(jour_lbl, ""),
                                'nuit_name': dict_agents.get(nuit_lbl, ""),
                                'nuit_comp': dict_comp_string.get(nuit_lbl, "")
                            })
                        st.divider()
                        export_data.append(veh_data)

        afficher_bloc_engin(vsav_uniques, '<div class="vsav-title">🚑 VÉHICULES DE SECOURS AUX VICTIMES (VSAV)</div>')
        
        groupes_par_taille = {}
        for agres in autres_uniques:
            nb_postes = len(df_garde[df_garde['Agrès'] == agres])
            if nb_postes not in groupes_par_taille: groupes_par_taille[nb_postes] = []
            groupes_par_taille[nb_postes].append(agres)

        ordre_tailles = [4, 6, 3, 2, 5]
        for taille in sorted(groupes_par_taille.keys(), reverse=True):
            if taille not in ordre_tailles: ordre_tailles.append(taille)

        for nb_postes in ordre_tailles:
            if nb_postes in groupes_par_taille:
                afficher_bloc_engin(groupes_par_taille[nb_postes], f'<div class="section-title">Équipages à {nb_postes} postes</div>')

        with st.sidebar:
            st.markdown("### 📥 Actions")
            
            if os.path.exists("modele.xlsx"):
                try:
                    wb = openpyxl.load_workbook("modele.xlsx")
                    ws = wb['BILLET'] if 'BILLET' in wb.sheetnames else wb.active
                    
                    found_anchors = []
                    for c in range(1, 30):
                        for r in range(1, 150):
                            val = ws.cell(row=r, column=c).value
                            if isinstance(val, str):
                                val_u = val.strip().upper()
                                for base in KNOWN_BASES:
                                    if val_u == base or val_u.startswith(base + " ") or val_u.startswith(base + "\n"):
                                        found_anchors.append({'r': r, 'c': c, 'base': base})
                                        break
                    
                    used_anchors = set()
                    for veh in export_data:
                        b_name = veh['base_name']
                        matched = None
                        
                        for idx, anc in enumerate(found_anchors):
                            if idx not in used_anchors and anc['base'] == b_name:
                                matched = anc
                                used_anchors.add(idx)
                                break
                        
                        if matched:
                            if veh['custom_name'].strip():
                                ws.cell(row=matched['r'], column=matched['c'], value=veh['custom_name'])
                                
                            curr_r = matched['r'] + 1
                            c = matched['c']
                            
                            for role in veh['roles']:
                                func = role['fonction']
                                found_r = None
                                for search_r in range(curr_r, curr_r + 20): 
                                    val = ws.cell(row=search_r, column=c).value
                                    if isinstance(val, str) and val.strip().upper() == func.upper():
                                        found_r = search_r
                                        break
                                        
                                if found_r:
                                    ws.cell(row=found_r, column=c+1, value=role['jour_name'])
                                    ws.cell(row=found_r+1, column=c+1, value=role['jour_comp']) 
                                    
                                    ws.cell(row=found_r, column=c+2, value=role['nuit_name'])
                                    ws.cell(row=found_r+1, column=c+2, value=role['nuit_comp']) 
                                    
                                    curr_r = found_r + 1
                    
                    output = BytesIO()
                    wb.save(output)
                    
                    st.download_button(
                        label="📥 Télécharger la Feuille Parfaite", 
                        data=output.getvalue(), 
                        file_name="Feuille_Garde_Jour_Nuit.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                        use_container_width=True
                    )
                except Exception as e:
                    st.error(f"Erreur d'édition Excel : {e}")
            else:
                st.error("❌ Fichier 'modele.xlsx' introuvable sur le serveur.")

except Exception as e:
    st.error(f"⚠️ Erreur Globale : {e}")
